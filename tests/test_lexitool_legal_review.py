import shutil
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt
from lxml import etree

sys.path.insert(0, str((Path(__file__).resolve().parents[1] / "vendor" / "lexitool").resolve()))

from lexitool.diff import summary
from lexitool.edit_ops import replace_text
from lexitool.markup import lex_read
from lexitool.openxml_opc import (
    CT_NS,
    PKG_REL_NS,
    append_relationship,
    ensure_default_content_type,
    ensure_override_content_type,
    rels_member_for_part,
    resolve_target_path,
)
from lexitool.openxml_runmap import render_paragraph, replace_span
from hermes_cli.project_commands import (
    edit_verification_record,
    legal_handoff_record,
    legal_harness_migrate,
    legal_harness_workflow,
    legal_review_plan,
    legal_scorecard,
    lex_convention_profile,
)
from tools.lexitool_tool import _handle_comment, _handle_edit, _handle_format, _handle_revision_guard, _handle_scan


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = f"{{{W_NS}}}"


def _inject_tracked_change(path: Path, paragraph_index: int, inserted: str, deleted: str = "") -> None:
    tmp = path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        doc_xml = zin.read("word/document.xml")
        root = etree.fromstring(doc_xml)
        paras = list(root.iter(f"{W}p"))
        para = paras[paragraph_index]

        if deleted:
            del_el = etree.Element(f"{W}del")
            del_el.set(f"{W}id", "42")
            del_el.set(f"{W}author", "JT")
            run = etree.SubElement(del_el, f"{W}r")
            t = etree.SubElement(run, f"{W}delText")
            t.text = deleted
            para.append(del_el)

        ins_el = etree.Element(f"{W}ins")
        ins_el.set(f"{W}id", "43")
        ins_el.set(f"{W}author", "JT")
        run = etree.SubElement(ins_el, f"{W}r")
        t = etree.SubElement(run, f"{W}t")
        t.text = inserted
        para.append(ins_el)

        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(item, data)
    shutil.move(tmp, path)


def test_legal_structure_detects_chinese_legal_headings(tmp_path):
    path = tmp_path / "support.docx"
    doc = Document()
    doc.add_paragraph("支持函")
    doc.add_paragraph("一、本公司承诺")
    doc.add_paragraph("本公司将维持对项目运营主体的控制权。")
    doc.add_paragraph("二、本函义务的履行")
    doc.save(path)

    outline = lex_read(str(path), mode="legal_structure")

    assert "§2 [L1] 一、本公司承诺" in outline
    assert "§4 [L1] 二、本函义务的履行" in outline


def test_review_digest_reports_tc_hotspots_and_next_reads(tmp_path):
    path = tmp_path / "review.docx"
    doc = Document()
    doc.add_paragraph("一、本公司承诺")
    doc.add_paragraph("本公司对项目运营主体提供支持。")
    doc.save(path)
    _inject_tracked_change(path, 1, "本公司对此承担不可撤销的流动性支持义务。", "本公司不承担支持义务。")

    digest = lex_read(str(path), mode="review")

    assert "[legal-review-digest]" in digest
    assert "Tracked-change paragraphs: 1" in digest
    assert "obligations/support" in digest
    assert "lex_read paras=[2]" in digest


def test_diff_summary_classifies_legal_revision_intent(tmp_path):
    original = tmp_path / "original.docx"
    revised = tmp_path / "revised.docx"
    doc = Document()
    doc.add_paragraph("本函不构成担保和债务承担。")
    doc.save(original)
    doc = Document()
    doc.add_paragraph("本函不构成担保或债务承担，但不影响本公司按照本函承诺对项目运营主体提供支持。")
    doc.save(revised)

    result = summary(str(original), str(revised))

    assert result["ok"] is True
    assert result["changes_total"] == 1
    assert "liability/security" in result["changes"][0]["categories"]
    assert "obligations/support" in result["changes"][0]["categories"]


def test_replace_all_requires_confirmation_for_high_risk_bulk_legal_edits(tmp_path):
    path = tmp_path / "bulk.docx"
    doc = Document()
    for i in range(10):
        doc.add_paragraph(f"本公司第{i}项承诺。")
    doc.save(path)

    result = _handle_edit({
        "path": str(path),
        "op": "replace_all",
        "targets": list(range(1, 10)),
        "old_text": "本",
        "new_text": "该",
    })

    assert "LEGAL_BULK_REPLACE_REVIEW_REQUIRED" in result


def test_direct_replace_preserves_unmatched_runs_when_match_crosses_runs(tmp_path):
    path = tmp_path / "share-mortgage.docx"
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("The ")
    p.add_run("Security").bold = True
    p.add_run(" Trustee").italic = True
    p.add_run(" may act.")
    doc.save(path)

    result = replace_text(str(path), 0, "Security Trustee", "Chargee", tc=False)

    assert result.ok is True

    with zipfile.ZipFile(path, "r") as zf:
        root = etree.fromstring(zf.read("word/document.xml"))
    para = next(root.iter(f"{W}p"))
    text_nodes = list(para.iter(f"{W}t"))
    texts = [t.text or "" for t in text_nodes]

    assert "".join(texts) == "The Chargee may act."
    assert texts == ["The ", "Chargee", "", " may act."]
    assert para.find(f".//{W}b") is not None
    assert para.find(f".//{W}i") is not None


def test_tc_replace_marks_only_matched_text_and_preserves_run_format(tmp_path):
    path = tmp_path / "share-mortgage-tc.docx"
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("The ")
    p.add_run("Security").bold = True
    p.add_run(" Trustee").italic = True
    p.add_run(" may act.")
    doc.save(path)

    result = replace_text(str(path), 0, "Security Trustee", "Chargee", tc=True, author="JT")

    assert result.ok is True

    with zipfile.ZipFile(path, "r") as zf:
        root = etree.fromstring(zf.read("word/document.xml"))
    para = next(root.iter(f"{W}p"))
    visible_text = "".join(t.text or "" for t in para.iter(f"{W}t"))
    deleted_text = "".join(t.text or "" for t in para.iter(f"{W}delText"))
    del_nodes = list(para.iter(f"{W}del"))
    ins_nodes = list(para.iter(f"{W}ins"))

    assert visible_text == "The Chargee may act."
    assert deleted_text == "Security Trustee"
    assert len(del_nodes) == 2
    assert len(ins_nodes) == 1
    assert del_nodes[0].find(f".//{W}b") is not None
    assert del_nodes[1].find(f".//{W}i") is not None
    assert ins_nodes[0].find(f".//{W}i") is not None


def test_openxml_runmap_renders_tabs_and_rejects_noneditable_span():
    para = etree.fromstring(
        f"""
        <w:p xmlns:w="{W_NS}">
          <w:r><w:t>A</w:t></w:r>
          <w:r><w:tab/></w:r>
          <w:r><w:t>B</w:t></w:r>
        </w:p>
        """.encode()
    )

    rendered = render_paragraph(para)

    assert rendered.text == "A\tB"
    assert [span.editable for span in rendered.spans] == [True, False, True]
    assert replace_span(rendered, 0, 3, "C") is False


def test_openxml_runmap_replace_span_preserves_surrounding_text_nodes():
    para = etree.fromstring(
        f"""
        <w:p xmlns:w="{W_NS}">
          <w:r><w:t>AA</w:t></w:r>
          <w:r><w:t>BB</w:t></w:r>
          <w:r><w:t>CC</w:t></w:r>
        </w:p>
        """.encode()
    )
    rendered = render_paragraph(para)

    assert replace_span(rendered, 1, 5, "X") is True

    texts = [t.text or "" for t in para.iter(f"{W}t")]
    assert texts == ["AX", "", "C"]
    assert "".join(texts) == "AXC"


def test_openxml_opc_helpers_manage_relationships_and_content_types():
    assert rels_member_for_part("word/document.xml") == "word/_rels/document.xml.rels"
    assert resolve_target_path("word/document.xml", "header1.xml") == "word/header1.xml"
    assert resolve_target_path("word/document.xml", "/custom/item1.xml") == "custom/item1.xml"

    rels = etree.fromstring(
        f'<Relationships xmlns="{PKG_REL_NS}"><Relationship Id="rId2" Type="old" Target="x"/></Relationships>'
    )
    rid = append_relationship(
        rels,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        "https://example.com",
        target_mode="External",
    )

    assert rid == "rId3"
    added = rels.findall(f"{{{PKG_REL_NS}}}Relationship")[-1]
    assert added.get("TargetMode") == "External"
    assert added.get("Target") == "https://example.com"

    types = etree.fromstring(f'<Types xmlns="{CT_NS}"><Default Extension="xml" ContentType="old/type"/></Types>')
    assert ensure_default_content_type(types, "xml", "application/xml") is True
    assert ensure_default_content_type(types, ".png", "image/png") is True
    assert ensure_default_content_type(types, "png", "image/png") is False
    assert ensure_override_content_type(types, "/word/comments.xml", "application/comments+xml") is True
    assert ensure_override_content_type(types, "word/comments.xml", "application/comments+xml") is False
    defaults = {el.get("Extension"): el.get("ContentType") for el in types.findall(f"{{{CT_NS}}}Default")}
    assert defaults == {"xml": "application/xml", "png": "image/png"}
    overrides = {el.get("PartName"): el.get("ContentType") for el in types.findall(f"{{{CT_NS}}}Override")}
    assert overrides == {"/word/comments.xml": "application/comments+xml"}


def test_lex_comment_add_precisely_anchors_cross_run_range(tmp_path):
    path = tmp_path / "comments.docx"
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("The ")
    p.add_run("Security").bold = True
    p.add_run(" Trustee").italic = True
    p.add_run(" may act.")
    doc.save(path)

    result = _handle_comment({
        "path": str(path),
        "op": "add",
        "para": 0,
        "range_start": 4,
        "range_end": 20,
        "author": "JT",
        "text": "Confirm this should be Chargee.",
    })

    assert '"ok": true' in result
    assert '"anchored_precisely": true' in result

    listed = _handle_comment({"path": str(path), "op": "list"})
    assert '"para": 0' in listed
    assert '"para_display": 1' in listed

    with zipfile.ZipFile(path, "r") as zf:
        doc_root = etree.fromstring(zf.read("word/document.xml"))
        rels_root = etree.fromstring(zf.read("word/_rels/document.xml.rels"))
        ct_root = etree.fromstring(zf.read("[Content_Types].xml"))

    para = next(doc_root.iter(f"{W}p"))
    child_tags = [etree.QName(child).localname for child in para]
    texts = ["".join(t.text or "" for t in child.iter(f"{W}t")) for child in para]

    assert "".join(t.text or "" for t in para.iter(f"{W}t")) == "The Security Trustee may act."
    assert child_tags == ["r", "commentRangeStart", "r", "r", "commentRangeEnd", "r", "r"]
    assert texts[0] == "The "
    assert texts[2] == "Security"
    assert texts[3] == " Trustee"
    assert texts[-1] == " may act."
    assert para.find(f".//{W}commentReference") is not None

    rels = rels_root.findall(f"{{{PKG_REL_NS}}}Relationship")
    assert any(rel.get("Type", "").endswith("/comments") and rel.get("Target") == "comments.xml" for rel in rels)
    overrides = ct_root.findall(f"{{{CT_NS}}}Override")
    assert any(ov.get("PartName") == "/word/comments.xml" for ov in overrides)


def test_insert_paragraphs_inherits_anchor_format_and_skips_empty_items(tmp_path):
    path = tmp_path / "insert-format.docx"
    doc = Document()
    anchor = doc.add_paragraph()
    anchor.paragraph_format.first_line_indent = Pt(24)
    run = anchor.add_run("锚点段落")
    run.font.size = Pt(16)
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    doc.add_paragraph("后续段落")
    doc.save(path)

    result = _handle_edit({
        "path": str(path),
        "op": "insert_paragraphs",
        "after_para": 0,
        "tc": False,
        "paragraphs": [
            {"text": "新增第一段"},
            {"text": ""},
            {"text": "新增第二段"},
        ],
    })

    assert '"ok": true' in result
    assert "skipped 1 empty paragraphs" in result

    readback = lex_read(str(path), paras=[1, 2, 3, 4], mode="full", show_tc=True, show_format=True)
    assert "§2 [indent:2ch][font:16pt]新增第一段[/font]" in readback
    assert "§3 [indent:2ch][font:16pt]新增第二段[/font]" in readback
    assert "[highlight:yellow]新增第一段" not in readback
    assert "§4 后续段落" in readback


def test_insert_paragraphs_compat_handles_legacy_insert_signature(monkeypatch, tmp_path):
    import lexitool.edit_ops as edit_ops

    path = tmp_path / "legacy-insert.docx"
    doc = Document()
    doc.add_paragraph("锚点")
    doc.save(path)
    captured = {}

    def legacy_insert(docx_path, after_para, paragraphs, *, tc=True, author="agent", font="宋体", sz=22.0, output=None):
        captured["paragraphs"] = paragraphs
        captured["kwargs"] = {"tc": tc, "author": author, "font": font, "sz": sz, "output": output}
        return edit_ops.EditResult(ok=True, message="legacy inserted", path=output or docx_path)

    monkeypatch.setattr(edit_ops, "insert_paragraph_block", legacy_insert)

    result = _handle_edit({
        "path": str(path),
        "op": "insert_paragraphs",
        "after_para": 0,
        "tc": True,
        "format": {"size": "16pt"},
        "inherit_format": True,
        "skip_empty": True,
        "paragraphs": [
            {"text": "新增段落"},
            {"text": ""},
        ],
    })

    assert '"ok": true' in result
    assert "skipped 1 empty paragraphs" in result
    assert captured["paragraphs"] == [{"text": "新增段落"}]
    assert captured["kwargs"]["tc"] is True
    assert captured["kwargs"]["output"] == str(path)


def test_lex_format_accepts_indent_units_and_sets_paragraph_run_defaults(tmp_path):
    path = tmp_path / "format-units.docx"
    doc = Document()
    doc.add_paragraph("第一段")
    doc.add_paragraph("需要格式化")
    doc.save(path)

    result = _handle_format({
        "path": str(path),
        "target": "§2",
        "properties": {"indent": "2.66667ch", "size": "16pt", "font": "宋体"},
    })

    assert '"ok": true' in result
    readback = lex_read(str(path), paras=[2], mode="full", show_tc=True, show_format=True)
    assert "§2 [indent:2.66667ch][font:宋体,16pt]需要格式化[/font]" in readback


def test_lex_edit_header_footer_replace_handles_cross_run_text(tmp_path):
    path = tmp_path / "header-footer.docx"
    doc = Document()
    section = doc.sections[0]
    header_para = section.header.paragraphs[0]
    header_para.add_run("Project ")
    header_para.add_run("Security").bold = True
    header_para.add_run(" Trustee").italic = True
    doc.add_paragraph("Body")
    doc.save(path)

    result = _handle_edit({
        "path": str(path),
        "op": "replace_header_footer",
        "kind": "header",
        "old_text": "Security Trustee",
        "new_text": "Chargee",
    })

    assert '"total_replacements": 1' in result

    with zipfile.ZipFile(path, "r") as zf:
        header_name = next(name for name in zf.namelist() if name.startswith("word/header") and name.endswith(".xml"))
        root = etree.fromstring(zf.read(header_name))
    texts = [t.text or "" for t in root.iter(f"{W}t")]

    assert "".join(texts) == "Project Chargee"
    assert texts == ["Project ", "Chargee", ""]
    assert root.find(f".//{W}b") is not None
    assert root.find(f".//{W}i") is not None


def test_lex_scan_returns_header_footer_targets(tmp_path):
    path = tmp_path / "header-scan.docx"
    doc = Document()
    section = doc.sections[0]
    header_para = section.header.paragraphs[0]
    header_para.add_run("Project ")
    header_para.add_run("Security").bold = True
    header_para.add_run(" Trustee").italic = True
    doc.add_paragraph("Body without target")
    doc.save(path)

    result = _handle_scan({
        "path": str(path),
        "query": "Security Trustee",
    })

    assert '"total_matches": 1' in result
    assert '"header_footer_targets"' in result
    assert '"kind": "header"' in result
    assert '"part_path": "word/header1.xml"' in result
    assert '"ref_types": ["default"]' in result


def test_revision_guard_blocks_unresolved_final_view_residuals(tmp_path):
    path = tmp_path / "share-mortgage.docx"
    doc = Document()
    doc.add_paragraph("“Secured Assets” means assets created in favour of the Security  Trustee.")
    doc.add_paragraph("The Chargee may enforce the Security.")
    doc.save(path)

    scan = _handle_scan({"path": str(path), "query": "Security Trustee"})
    assert '"total_matches": 1' in scan
    assert '"paragraph_targets": [1]' in scan

    failed = _handle_revision_guard(
        {
            "path": str(path),
            "required_absent": ["Security Trustee"],
            "required_present": ["Chargee"],
        }
    )
    assert '"ok": false' in failed
    assert '"failure_count": 1' in failed
    assert '"paragraph_targets": [1]' in failed

    passed = _handle_revision_guard(
        {
            "path": str(path),
            "required_absent": ["Security Agent"],
            "required_present": ["Chargee"],
        }
    )
    assert '"ok": true' in passed
    assert '"failure_count": 0' in passed


def test_legal_harness_primitives_persist_project_state(tmp_path):
    project = tmp_path / "matter"
    hp = project / ".hermes-project"
    hp.mkdir(parents=True)
    (hp / "project-meta.json").write_text(
        '{"name":"测试项目","client":"客户","goal":"测试"}',
        encoding="utf-8",
    )
    doc_path = project / "support.docx"
    doc = Document()
    doc.add_paragraph("支持函")
    doc.add_paragraph("一、本公司承诺")
    doc.add_paragraph("本函不构成担保或债务承担，但本公司承诺提供流动性支持。")
    doc.save(doc_path)

    migrated = legal_harness_migrate(project_dirs=[str(project)])
    profile = lex_convention_profile(str(doc_path), project_dir=str(project))
    plan = legal_review_plan(str(doc_path), project_dir=str(project))
    record = edit_verification_record(
        str(project),
        str(doc_path),
        target="§3",
        edit_summary="确认非担保表述和支持义务并存",
        before_text="本函不构成担保。",
        after_text="本函不构成担保或债务承担，但本公司承诺提供流动性支持。",
        checks=["lex_read readback", "liability/security", "obligations/support"],
        status="passed",
    )

    assert migrated["migrated_count"] == 1
    assert profile["ok"] is True
    assert profile["profile"]["xref_convention"]["summary"] in {"unknown", "第X条"}
    assert plan["ok"] is True
    assert plan["plan"]["steps"]
    assert record["ok"] is True
    assert (hp / "convention-profiles.json").exists()
    assert (hp / "legal-review-plans.json").exists()
    assert (hp / "edit-verification-records.json").exists()
    assert "Edit verification passed" in (hp / "project-context.md").read_text(encoding="utf-8")


def test_legal_harness_workflows_handoffs_and_scorecard(tmp_path):
    project = tmp_path / "matter"
    hp = project / ".hermes-project"
    hp.mkdir(parents=True)
    (hp / "project-meta.json").write_text('{"name":"测试项目"}', encoding="utf-8")
    doc_path = project / "support.docx"
    doc = Document()
    doc.add_paragraph("支持函")
    doc.add_paragraph("一、本公司承诺")
    doc.add_paragraph("本函不构成担保或债务承担，但本公司承诺提供流动性支持。")
    doc.save(doc_path)

    migrated = legal_harness_migrate(project_dirs=[str(project)])
    workflows = legal_harness_workflow(str(project))
    workflow = legal_harness_workflow(str(project), action="get", workflow_id="contract_revision")

    assert migrated["ok"] is True
    assert (hp / "workflows" / "contract_revision.yaml").exists()
    assert workflows["count"] >= 4
    assert workflow["workflow"]["nodes"]

    failed_handoff = legal_handoff_record(
        str(project),
        workflow_id="contract_revision",
        node_id="revise",
        handoff={"status": "completed", "evidence": {}},
    )
    assert failed_handoff["ok"] is False
    assert "modified_files" in failed_handoff["error"]

    good_handoff = legal_handoff_record(
        str(project),
        workflow_id="contract_revision",
        node_id="revise",
        handoff={
            "status": "completed",
            "modified_files": [str(doc_path)],
            "modified_locations": ["§3"],
            "guards": {"revision_guard": {"ok": True}},
            "evidence": {"readback": "§3 checked"},
        },
    )
    assert good_handoff["ok"] is True
    run_id = good_handoff["record"]["run_id"]
    assert (hp / "harness-runs" / run_id / "nodes" / "revise.json").exists()

    failing_scorecard = legal_scorecard(
        str(project),
        document_path=str(doc_path),
        workflow_id="contract_revision",
        run_id=run_id,
    )
    assert failing_scorecard["ok"] is False
    assert any(item["check"] == "convention_profile" for item in failing_scorecard["failures"])

    lex_convention_profile(str(doc_path), project_dir=str(project))
    legal_review_plan(str(doc_path), project_dir=str(project))
    edit_verification_record(
        str(project),
        str(doc_path),
        target="§3",
        edit_summary="确认支持义务",
        before_text="本函不构成担保。",
        after_text="本函不构成担保或债务承担，但本公司承诺提供流动性支持。",
        checks=["lex_read readback", "lex_revision_guard"],
        status="passed",
    )
    passing_scorecard = legal_scorecard(
        str(project),
        document_path=str(doc_path),
        workflow_id="contract_revision",
        run_id=run_id,
    )
    assert passing_scorecard["ok"] is True
    assert passing_scorecard["status"] == "passed"


def test_insert_text_inherits_paragraph_format(tmp_path):
    """insert_text(inherit_format=True) copies the first run's rPr from the paragraph."""
    from lexitool.edit_ops import insert_text

    path = tmp_path / "inherit.docx"
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run("现有文本")
    run.font.size = Pt(16)
    run.font.name = "Calibri"
    doc.save(path)

    result = insert_text(str(path), 0, "追加文本", inherit_format=True, tc=False)

    assert result.ok is True

    with zipfile.ZipFile(path, "r") as zf:
        root = etree.fromstring(zf.read("word/document.xml"))
    para = next(root.iter(f"{W}p"))
    runs = list(para.iter(f"{W}r"))
    last_run = runs[-1]
    rPr = last_run.find(f"{W}rPr")
    assert rPr is not None
    sz = rPr.find(f"{W}sz")
    assert sz is not None
    assert sz.get(f"{W}val") == "32"  # 16pt × 2
    rFonts = rPr.find(f"{W}rFonts")
    assert rFonts is not None
    assert rFonts.get(f"{W}ascii") == "Calibri"


def test_insert_text_explicit_format_overrides_inherited(tmp_path):
    """Explicit font/size kwargs override inherited paragraph formatting."""
    from lexitool.edit_ops import insert_text

    path = tmp_path / "override.docx"
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run("现有文本")
    run.font.size = Pt(16)
    run.font.name = "Calibri"
    doc.save(path)

    result = insert_text(str(path), 0, "追加文本",
                         font="Times New Roman", font_size=14.0,
                         inherit_format=True, tc=False)

    assert result.ok is True

    with zipfile.ZipFile(path, "r") as zf:
        root = etree.fromstring(zf.read("word/document.xml"))
    para = next(root.iter(f"{W}p"))
    last_run = list(para.iter(f"{W}r"))[-1]
    rPr = last_run.find(f"{W}rPr")
    sz = rPr.find(f"{W}sz")
    assert sz.get(f"{W}val") == "28"  # 14pt × 2
    rFonts = rPr.find(f"{W}rFonts")
    assert rFonts.get(f"{W}ascii") == "Times New Roman"


def test_find_and_replace_all_tc_mode_produces_del_and_ins(tmp_path):
    """find_and_replace_all(tc=True) wraps changes in w:del / w:ins."""
    from lexitool.edit_ops import find_and_replace_all

    path = tmp_path / "tc-replace-all.docx"
    doc = Document()
    doc.add_paragraph("本公司承诺提供支持。")
    doc.add_paragraph("本公司不承担担保义务。")
    doc.save(path)

    result = find_and_replace_all(str(path), "本公司", "甲方", tc=True, author="JT")

    assert result["ok"] is True
    assert result["match_count"] == 2

    with zipfile.ZipFile(path, "r") as zf:
        root = etree.fromstring(zf.read("word/document.xml"))
    del_elements = list(root.iter(f"{W}del"))
    ins_elements = list(root.iter(f"{W}ins"))
    assert len(del_elements) >= 2
    assert len(ins_elements) >= 2

    del_texts = "".join(t.text or "" for t in root.iter(f"{W}delText"))
    ins_texts = "".join(t.text or "" for t in root.iter(f"{W}t"))
    assert "本公司" in del_texts
    assert "甲方" in ins_texts


def test_reload_purges_vendored_modules():
    """reload_lexitool_tools() purges lexitool.* from sys.modules before re-import."""
    import sys as _sys
    from tools.lexitool_tool import reload_lexitool_tools

    # Ensure a vendored module is present before reload
    import lexitool.edit_ops  # noqa: F811
    assert "lexitool.edit_ops" in _sys.modules

    result = reload_lexitool_tools()

    assert result["vendored_modules_purged"] > 0
    # After purge, vendored modules are re-imported lazily on first use.
    import importlib
    eo = importlib.import_module("lexitool.edit_ops")
    assert eo is not None


def test_legal_harness_tools_are_exposed_in_lex_toolsets():
    from toolsets import resolve_toolset

    lex_tools = set(resolve_toolset("lexitool"))
    coordinator_tools = set(resolve_toolset("lex-docx-coordinator"))
    worker_tools = set(resolve_toolset("lex-docx-worker"))

    for name in {"legal_harness_workflow", "legal_handoff_record", "legal_scorecard"}:
        assert name in lex_tools
        assert name in coordinator_tools
        assert name in worker_tools
