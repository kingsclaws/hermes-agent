import sys
from pathlib import Path

from docx import Document
from lxml import etree

sys.path.insert(0, str((Path(__file__).resolve().parents[1] / "vendor" / "lexitool").resolve()))

from lexitool.fields import list_fields
from lexitool.xref import audit_documents, convert_static_refs


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = etree.Element(f"{W}bookmarkStart")
    start.set(f"{W}id", str(bookmark_id))
    start.set(f"{W}name", name)
    end = etree.Element(f"{W}bookmarkEnd")
    end.set(f"{W}id", str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def test_convert_static_refs_routes_adjacent_clauses_exactly(tmp_path):
    path = tmp_path / "xref.docx"
    doc = Document()
    p1 = doc.add_paragraph("17.1 财务承诺")
    p1.style = "Heading 2"
    _add_bookmark(p1, "_Ref171", 171)
    p2 = doc.add_paragraph("17.2 财务测算")
    p2.style = "Heading 2"
    _add_bookmark(p2, "_Ref172", 172)
    doc.add_paragraph("融资方应遵守本协议第17.2条的约定，并另见第17.1条。")
    doc.save(path)

    dry = convert_static_refs(str(path), dry_run=True)
    refs = {item["clause"]: item["bookmark"] for item in dry["details"]}

    assert refs["17.1"] == "_Ref171"
    assert refs["17.2"] == "_Ref172"

    done = convert_static_refs(str(path), dry_run=False)
    fields = list_fields(str(path))
    bookmarks = [item.get("bookmark") for item in fields if item.get("field_type") == "REF"]

    assert done["converted"] == 2
    assert "_Ref171" in bookmarks
    assert "_Ref172" in bookmarks


def test_audit_documents_handles_single_doc_projects(tmp_path):
    path = tmp_path / "single.docx"
    doc = Document()
    doc.add_paragraph("1 定义").style = "Heading 1"
    doc.add_paragraph("本协议第9条不存在。")
    doc.save(path)

    result = audit_documents([str(path)])

    assert result["docs_scanned"] == 1
    assert "internal_dead" in result["summary"]
